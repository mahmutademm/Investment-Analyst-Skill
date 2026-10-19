"""Deep-dive workbook: <TICKER>_DeepDive_<date>.xlsx

Every ratio, gate test, DCF line and decision number is a live formula pointing back to the
Financials sheet (SEC XBRL values with their concept tags on the 'XBRL Sources' sheet).
"""
from __future__ import annotations

import math
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter as L

from iat import ratios, store, valuation
from iat.company import GATE_SPEC, frame_from_json, thr
from iat.config import Workspace
from iat.excel import common as C
from iat.fundamentals import FIELDS

MONEY_FIELDS_2DP = {"eps_diluted", "price_at_period_end"}


def _f(x):
    try:
        x = float(x)
        return x if math.isfinite(x) else None
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------
# Financials + Ratios
# ---------------------------------------------------------------------------

def sheet_financials(wb, d) -> tuple[dict, list[str], list[str]]:
    fr = frame_from_json(d["fundamentals"])
    src = frame_from_json(d["fundamentals_source"])
    ws = wb.create_sheet("Financials")
    cols = list(fr.columns)
    r = C.title(ws, f"{d['ticker']} standardized financials ({d['source'].upper()})",
                "Values in USD (displayed in millions). Concept used for each value: 'XBRL Sources' sheet. "
                "Blank = not reported/tagged. You may type over a value after checking the 10-K.")
    hdr_row = r
    ends = d["fundamentals_meta"].get("period_ends", {})
    C.header(ws, hdr_row, ["Key", "Line item", *[f"{c}" + (f" ({ends.get('TTM')})" if c == "TTM" else "") for c in cols],
                           "Source of latest value"])
    letters = [L(3 + i) for i in range(len(cols))]
    rows = {}
    for i, key in enumerate(FIELDS, start=1):
        rr = hdr_row + i
        rows[key] = rr
        C.put(ws, rr, 1, key).font = C.MUTED
        C.put(ws, rr, 2, FIELDS[key])
        fmt = C.NUM2 if key in MONEY_FIELDS_2DP else C.MILL
        for j, c in enumerate(cols):
            v = fr.loc[key, c] if key in fr.index else None
            C.put(ws, rr, 3 + j, v, fmt)
        C.put(ws, rr, 3 + len(cols), src.loc[key, cols[-1]] if key in src.index else "").font = C.MUTED
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 36
    for le in letters:
        ws.column_dimensions[le].width = 14
    ws.column_dimensions[L(3 + len(cols))].width = 70
    C.freeze(ws, f"C{hdr_row + 1}")

    ws2 = wb.create_sheet("XBRL Sources")
    r2 = C.title(ws2, "Where each number came from", "XBRL concept, form and accession number per cell. "
                                                     "Search the accession number on sec.gov to open the filing.")
    C.header(ws2, r2, ["Key", *cols])
    for i, key in enumerate(FIELDS, start=1):
        C.put(ws2, r2 + i, 1, key)
        for j, c in enumerate(cols):
            C.put(ws2, r2 + i, 2 + j, src.loc[key, c] if key in src.index else "").font = C.MUTED
    ws2.column_dimensions["A"].width = 22
    for j in range(len(cols)):
        ws2.column_dimensions[L(2 + j)].width = 48
    return rows, cols, letters


def sheet_ratios(wb, d, fin_rows, cols, letters) -> dict:
    ws = wb.create_sheet("Ratios")
    r = C.title(ws, "Ratios (live formulas on the Financials sheet)",
                "Growth/average-based ratios are blank in the TTM column on purpose (TTM vs FY is not a year-over-year change).")
    hdr = r
    C.header(ws, hdr, ["Key", "Ratio", *cols, "What to look for"])
    metric_rows: dict[str, int] = {}
    layout = []
    rr = hdr + 1
    for g in ratios.groups():
        layout.append(("group", g, rr))
        rr += 1
        for rt in [x for x in ratios.R if x.group == g]:
            metric_rows[rt.key] = rr
            layout.append(("ratio", rt, rr))
            rr += 1
    ttm = {i for i, c in enumerate(cols) if c == "TTM"}
    for kind, obj, row in layout:
        if kind == "group":
            ws.cell(row=row, column=2, value=obj).font = C.SECTION_FONT
            continue
        C.put(ws, row, 1, obj.key).font = C.MUTED
        C.put(ws, row, 2, obj.label)
        for i in range(len(cols)):
            ctx = ratios.XlCtx(fin_rows, metric_rows, letters, i, ttm)
            C.put(ws, row, 3 + i, ratios.excel_formula(obj, ctx), obj.fmt)
        C.put(ws, row, 3 + len(cols), obj.note, wrap=True).font = C.MUTED
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 24, 46
    for le in letters:
        ws.column_dimensions[le].width = 12
    ws.column_dimensions[L(3 + len(cols))].width = 60
    C.freeze(ws, f"C{hdr + 1}")
    return metric_rows


# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------

SCEN_DRIVERS = [
    ("g1", "Stage-1 revenue growth (per year)", C.PCT),
    ("n1", "Stage-1 years (then fades to terminal by year 10)", "0"),
    ("g_terminal", "Terminal growth", C.PCT),
    ("margin_target", "Target EBIT margin", C.PCT),
    ("margin_years", "Years to reach target margin", "0"),
    ("sales_to_capital", "Sales-to-capital (revenue added per $ reinvested)", "0.00"),
    ("ronic", "Return on new capital in perpetuity", C.PCT),
    ("wacc_adj", "WACC adjustment for this scenario", "0.00%"),
    ("probability", "Probability", "0%"),
]


def sheet_inputs(wb, d, notes, fin_rows, metric_rows, cols, letters) -> dict:
    ws = wb.create_sheet("Inputs")
    m = d["market"]
    lvl = letters[-1]
    r = C.title(ws, "Inputs and assumptions", "Blue cells are yours to change. Every valuation formula reads from here.")
    refs: dict[str, str] = {}

    def line(row, key, label, value, fmt, inp, note="", define=True):
        C.put(ws, row, 2, label)
        C.put(ws, row, 3, value, fmt, inp=inp)
        C.put(ws, row, 4, note, wrap=True).font = C.MUTED
        if define:
            C.name(wb, key, "Inputs", row, 3)
        refs[key] = key
        return row + 1

    r = C.section(ws, r, "Market and capital structure")
    r = line(r, "Price", "Share price", m["price"], C.NUM2, True, "Yahoo quote at build time")
    r = line(r, "Shares", "Shares outstanding", m["shares"], C.NUM0, True, "SEC cover page (dei) or diluted weighted shares")
    r = line(r, "MarketCap", "Market value of equity", "=Price*Shares", C.MILL, False)
    r = line(r, "Beta", "Beta", m["beta"], "0.00", True, "Yahoo (5y monthly). Consider a peer median or Blume-adjusted beta")
    r = line(r, "Rf", "Risk-free rate (10Y Treasury)", m["rf"], "0.00%", True, m.get("rf_source", ""))
    r = line(r, "ERP", "Equity risk premium", m["erp"], "0.00%", True, "settings.valuation.equity_risk_premium; check Damodaran monthly")
    r = line(r, "Ke", "Cost of equity (CAPM)", "=Rf+Beta*ERP", "0.00%", False)
    r = line(r, "Kd", "Pre-tax cost of debt", m["pre_tax_cost_of_debt"], "0.00%", True,
             "Interest / avg debt, or rf + 2% fallback. Compare with bond yields / rating")
    r = line(r, "TaxRate", "Tax rate", m["tax_rate"], "0.0%", True, "3-year average effective rate (clipped 0-35%)")
    r = line(r, "KdAfterTax", "After-tax cost of debt", "=Kd*(1-TaxRate)", "0.00%", False)
    r = line(r, "Debt", "Debt incl. leases", f"=N(Ratios!{lvl}{metric_rows['total_debt_incl_leases']})", C.MILL, False,
             "Latest column of Ratios")
    r = line(r, "WACC", "WACC", "=MarketCap/(MarketCap+Debt)*Ke+Debt/(MarketCap+Debt)*KdAfterTax", "0.00%", False,
             "Market-value weights; debt at book")
    r = line(r, "NetDebt", "Net debt incl. leases", f"=N(Ratios!{lvl}{metric_rows['net_debt']})", C.MILL, False)
    r = line(r, "Minority", "Non-controlling interest", f"=N(Financials!{lvl}{fin_rows['minority_interest']})", C.MILL, False)
    r = line(r, "RevBase", "Base revenue (latest column)", f"=Financials!{lvl}{fin_rows['revenue']}", C.MILL, False)
    r = line(r, "MarginBase", "Current EBIT margin", f"=Ratios!{lvl}{metric_rows['operating_margin']}", C.PCT, False)
    r = line(r, "FwdEPS", "Consensus forward EPS", m.get("forward_eps"), C.NUM2, True, "Yahoo")
    r = line(r, "ATR", "ATR(14), daily", m.get("atr14"), C.NUM2, True, "Average true range for stop sizing")
    r += 1
    r = C.section(ws, r, "Portfolio rules (from settings.yaml)")
    p = wb._iat_settings["portfolio"]
    v = wb._iat_settings["valuation"]
    r = line(r, "Portfolio", "Portfolio value", p.get("value"), C.NUM0, True, "Set portfolio.value in settings.yaml")
    r = line(r, "RiskPct", "Risk per position (% of portfolio)", (p.get("risk_per_position_pct") or 1) / 100, "0.00%", True)
    r = line(r, "MaxPct", "Max position (% of portfolio)", (p.get("max_position_pct") or 10) / 100, "0.0%", True)
    r = line(r, "ATRk", "Stop distance (x ATR)", v.get("atr_multiple"), "0.0", True)
    tax_st = p.get("short_term_tax_rate_pct")
    r = line(r, "STTax", "Short-term capital gains tax rate", (tax_st / 100) if tax_st else 0, "0.0%", True,
             "Gains held < 1 year are taxed as ordinary income in the US. 0 = ignore")
    r += 1
    r = C.section(ws, r, "Scenario drivers (DCF)")
    C.header(ws, r, ["", "Driver", "Bear", "Base", "Bull", "Why (AI evidence + your judgment)"], 1)
    r += 1
    scen = {s["name"]: s for s in d["scenarios"]}
    ov = (notes.get("assumptions") or {})
    why = ov.get("rationale", {}) if isinstance(ov.get("rationale"), dict) else {}
    scen_cells: dict[tuple[str, str], str] = {}
    for key, label, fmt in SCEN_DRIVERS:
        C.put(ws, r, 2, label)
        for j, sname in enumerate(("Bear", "Base", "Bull")):
            val = (ov.get(sname) or {}).get(key, scen.get(sname, {}).get(key))
            C.put(ws, r, 3 + j, val, fmt, inp=True)
            scen_cells[(sname, key)] = f"Inputs!${L(3 + j)}${r}"
        C.put(ws, r, 6, why.get(key, ""), wrap=True, ai=bool(why.get(key)))
        r += 1
    C.put(ws, r, 2, "Probabilities must sum to 100%")
    C.put(ws, r, 3, f"=IF(ABS(SUM(C{r - 1}:E{r - 1})-1)<0.001,\"OK\",\"FIX: sum is \"&TEXT(SUM(C{r - 1}:E{r - 1}),\"0%\"))")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 3, 48
    ws.column_dimensions["C"].width = ws.column_dimensions["D"].width = ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 80
    refs["scen"] = scen_cells  # type: ignore[assignment]
    return refs


# ---------------------------------------------------------------------------
# Valuation
# ---------------------------------------------------------------------------

def sheet_valuation(wb, d, refs, notes, metric_rows, fin_rows, letters) -> dict:
    ws = wb.create_sheet("Valuation")
    sc = refs["scen"]
    r = C.title(ws, "Intrinsic value", "FCFF DCF (3-stage: stage-1 growth, linear fade, terminal), three scenarios, "
                                        "plus DDM / financials / multiples cross-checks.")
    out = {"vps": {}, "blocks": {}}
    for sname in ("Bear", "Base", "Bull"):
        g = lambda k: sc[(sname, k)]  # noqa: E731
        r = C.section(ws, r, f"{sname} scenario")
        hdr = r
        C.header(ws, hdr, ["", "Year", "0", *[str(t) for t in range(1, 11)], "Terminal"], 1)
        for t in range(1, 11):
            ws.cell(row=hdr, column=3 + t, value=t)
        rows = {k: hdr + i for i, k in enumerate(["growth", "revenue", "margin", "ebit", "nopat", "reinv", "fcff", "df", "pv"], 1)}
        labels = {"growth": "Revenue growth", "revenue": "Revenue", "margin": "EBIT margin", "ebit": "EBIT",
                  "nopat": "NOPAT = EBIT x (1 - tax)", "reinv": "Reinvestment", "fcff": "FCFF", "df": "Discount factor",
                  "pv": "PV of FCFF"}
        for k, rr in rows.items():
            C.put(ws, rr, 2, labels[k])
        wcell = f"$C${hdr + 11}"
        for t in range(0, 12):
            col = 3 + t
            le, prev = L(col), L(col - 1)
            tt = f"{le}${hdr}"
            if t == 0:
                C.put(ws, rows["revenue"], col, "=RevBase", C.MILL)
                C.put(ws, rows["margin"], col, "=MarginBase", C.PCT)
                C.put(ws, rows["ebit"], col, f"={le}{rows['revenue']}*{le}{rows['margin']}", C.MILL)
                continue
            if t <= 10:
                C.put(ws, rows["growth"], col,
                      f"=IF({tt}<={g('n1')},{g('g1')},{g('g1')}-({g('g1')}-{g('g_terminal')})*({tt}-{g('n1')})/(10-{g('n1')}))", C.PCT)
                C.put(ws, rows["margin"], col, f"=MarginBase+({g('margin_target')}-MarginBase)*MIN({tt},{g('margin_years')})/{g('margin_years')}", C.PCT)
                C.put(ws, rows["reinv"], col, f"=({le}{rows['revenue']}-{prev}{rows['revenue']})/{g('sales_to_capital')}", C.MILL)
                C.put(ws, rows["df"], col, f"=1/(1+{wcell})^{tt}", "0.0000")
                C.put(ws, rows["pv"], col, f"={le}{rows['fcff']}*{le}{rows['df']}", C.MILL)
            else:  # terminal column
                C.put(ws, rows["growth"], col, f"={g('g_terminal')}", C.PCT)
                C.put(ws, rows["margin"], col, f"={g('margin_target')}", C.PCT)
                C.put(ws, rows["reinv"], col, f"={le}{rows['nopat']}*{g('g_terminal')}/{g('ronic')}", C.MILL)
            C.put(ws, rows["revenue"], col, f"={prev}{rows['revenue']}*(1+{le}{rows['growth']})", C.MILL)
            C.put(ws, rows["ebit"], col, f"={le}{rows['revenue']}*{le}{rows['margin']}", C.MILL)
            C.put(ws, rows["nopat"], col, f"={le}{rows['ebit']}*(1-TaxRate)", C.MILL)
            C.put(ws, rows["fcff"], col, f"={le}{rows['nopat']}-{le}{rows['reinv']}", C.MILL)
        s0 = hdr + 11
        items = [
            ("Scenario WACC", f"=WACC+{g('wacc_adj')}", "0.00%"),
            ("Sum of PV(FCFF), years 1-10", f"=SUM(D{rows['pv']}:M{rows['pv']})", C.MILL),
            ("Terminal value (FCFF year 11 / (WACC - g))", f"=N{rows['fcff']}/(C{s0}-{g('g_terminal')})", C.MILL),
            ("PV of terminal value", f"=C{s0 + 2}/(1+C{s0})^10", C.MILL),
            ("Enterprise value", f"=C{s0 + 1}+C{s0 + 3}", C.MILL),
            ("Less: net debt incl. leases", "=NetDebt", C.MILL),
            ("Less: non-controlling interest", "=Minority", C.MILL),
            ("Equity value", f"=C{s0 + 4}-C{s0 + 5}-C{s0 + 6}", C.MILL),
            ("Value per share", f"=C{s0 + 7}/Shares", C.NUM2),
            ("Upside vs price", f"=C{s0 + 8}/Price-1", C.PCT),
            ("Terminal value share of EV", f"=C{s0 + 3}/C{s0 + 4}", C.PCT),
        ]
        for i, (lab, f, fmt) in enumerate(items):
            C.put(ws, s0 + i, 2, lab, bold=lab == "Value per share")
            C.put(ws, s0 + i, 3, f, fmt, bold=lab == "Value per share")
        tvw = wb._iat_settings["thresholds"].get("tv_share_warning", 0.7)
        C.put(ws, s0 + 10, 4, f"=IF(C{s0 + 10}>{tvw},\"WARNING: value depends mostly on the distant future\",\"OK\")")
        py = (d.get("valuation_python") or {}).get(sname, {})
        C.put(ws, s0 + 11, 2, "Python cross-check (default inputs at build time)").font = C.MUTED
        C.put(ws, s0 + 11, 3, py.get("per_share"), C.NUM2).font = C.MUTED
        out["vps"][sname] = f"Valuation!$C${s0 + 8}"
        out["blocks"][sname] = {"hdr": hdr, "rows": rows, "s0": s0}
        r = s0 + 13

    # probability-weighted
    r = C.section(ws, r, "Scenario summary")
    C.header(ws, r, ["", "Scenario", "Probability", "Value / share", "Upside"], 1)
    first = r + 1
    for i, sname in enumerate(("Bear", "Base", "Bull")):
        rr = first + i
        C.put(ws, rr, 2, sname)
        C.put(ws, rr, 3, f"={sc[(sname, 'probability')]}", "0%")
        C.put(ws, rr, 4, f"={out['vps'][sname]}", C.NUM2)
        C.put(ws, rr, 5, f"=D{rr}/Price-1", C.PCT)
    rr = first + 3
    C.put(ws, rr, 2, "Probability-weighted", bold=True)
    C.put(ws, rr, 4, f"=SUMPRODUCT(C{first}:C{first + 2},D{first}:D{first + 2})", C.NUM2, bold=True)
    C.put(ws, rr, 5, f"=D{rr}/Price-1", C.PCT, bold=True)
    C.name(wb, "PWValue", "Valuation", rr, 4)
    r = rr + 2

    # sensitivity (base)
    b = out["blocks"]["Base"]
    r = C.section(ws, r, "Sensitivity: base value per share, WACC (rows) x terminal growth (columns)")
    C.put(ws, r, 2, "Explicit-period cash flows held at the base case; terminal FCFF recomputed for each growth rate.").font = C.MUTED
    r += 1
    hdr = r
    C.put(ws, hdr, 2, "WACC \\ g", bold=True)
    offs = [-0.01, -0.005, 0.0, 0.005, 0.01]
    for j, o in enumerate(offs):
        C.put(ws, hdr, 3 + j, f"={sc[('Base', 'g_terminal')]}+({o})", "0.00%", bold=True)
    s0 = b["s0"]
    fc = f"$D${b['rows']['fcff']}:$M${b['rows']['fcff']}"
    rev10 = f"$M${b['rows']['revenue']}"
    for i, o in enumerate(offs):
        rr = hdr + 1 + i
        C.put(ws, rr, 2, f"=$C${s0}+({o})", "0.00%", bold=True)
        for j in range(len(offs)):
            gw, gg = f"$B{rr}", f"{L(3 + j)}${hdr}"
            f = (f"=IFERROR((NPV({gw},{fc})+({rev10}*(1+{gg})*{sc[('Base', 'margin_target')]}*(1-TaxRate))"
                 f"*(1-{gg}/{sc[('Base', 'ronic')]})/({gw}-{gg})/(1+{gw})^10-NetDebt-Minority)/Shares,\"\")")
            C.put(ws, rr, 3 + j, f, C.NUM2)
    C.heat(ws, f"C{hdr + 1}:G{hdr + 5}")
    r = hdr + 7

    # DDM
    r = C.section(ws, r, "Dividend discount model (use for stable dividend payers)")
    m = d["market"]
    ddm_rows = {}
    for lab, key, val, fmt, inp in [
        ("Current dividend per share (D0)", "D0", m.get("dps"), C.NUM2, True),
        ("Stage-1 dividend growth", "DDMg1", 0.05, C.PCT, True),
        ("Stage-1 years", "DDMn", 5, "0", True),
        ("Terminal growth", "DDMgT", f"={sc[('Base', 'g_terminal')]}", C.PCT, False),
        ("Required return (cost of equity)", "DDMr", "=Ke", "0.00%", False),
    ]:
        C.put(ws, r, 2, lab)
        C.put(ws, r, 3, val, fmt, inp=inp)
        ddm_rows[key] = f"$C${r}"
        r += 1
    C.put(ws, r, 2, "2-stage DDM value per share", bold=True)
    C.put(ws, r, 3, f"=IF(N({ddm_rows['D0']})<=0,\"n/a (no dividend)\",SUMPRODUCT({ddm_rows['D0']}*(1+{ddm_rows['DDMg1']})^ROW(INDIRECT(\"1:\"&{ddm_rows['DDMn']}))"
                    f"/(1+{ddm_rows['DDMr']})^ROW(INDIRECT(\"1:\"&{ddm_rows['DDMn']})))+{ddm_rows['D0']}*(1+{ddm_rows['DDMg1']})^{ddm_rows['DDMn']}"
                    f"*(1+{ddm_rows['DDMgT']})/({ddm_rows['DDMr']}-{ddm_rows['DDMgT']})/(1+{ddm_rows['DDMr']})^{ddm_rows['DDMn']})", C.NUM2, bold=True)
    out["ddm"] = f"Valuation!$C${r}"
    r += 1
    C.put(ws, r, 2, "Gordon growth value per share")
    C.put(ws, r, 3, f"=IF(N({ddm_rows['D0']})<=0,\"n/a\",{ddm_rows['D0']}*(1+{ddm_rows['DDMgT']})/({ddm_rows['DDMr']}-{ddm_rows['DDMgT']}))", C.NUM2)
    r += 2

    # Financials: justified P/B and residual income
    r = C.section(ws, r, "Financials: justified P/B and residual income (banks, insurers)")
    lvl = letters[-1]
    fin = {}
    for lab, key, val, fmt, inp in [
        ("Book value per share", "BVPS", f"=Financials!{lvl}{fin_rows['equity']}/Shares", C.NUM2, False),
        ("Sustainable ROE", "ROEs", m.get("roe_3y"), C.PCT, True),
        ("Long-run growth", "FINg", f"={sc[('Base', 'g_terminal')]}", C.PCT, False),
        ("Payout ratio", "Payout", 0.3, C.PCT, True),
    ]:
        C.put(ws, r, 2, lab)
        C.put(ws, r, 3, val, fmt, inp=inp)
        fin[key] = f"$C${r}"
        r += 1
    C.put(ws, r, 2, "Justified P/B = (ROE - g) / (Ke - g)")
    C.put(ws, r, 3, f"=IFERROR(({fin['ROEs']}-{fin['FINg']})/(Ke-{fin['FINg']}),\"\")", '0.00"x"')
    r += 1
    C.put(ws, r, 2, "Justified P/B value per share", bold=True)
    C.put(ws, r, 3, f"=IFERROR(C{r - 1}*{fin['BVPS']},\"\")", C.NUM2, bold=True)
    out["jpb"] = f"Valuation!$C${r}"
    r += 1
    C.put(ws, r, 2, "Residual income value per share (5y + terminal)")
    ri = (f"=IFERROR({fin['BVPS']}+SUMPRODUCT(({fin['ROEs']}-Ke)*{fin['BVPS']}*(1+{fin['ROEs']}*(1-{fin['Payout']}))^(ROW(INDIRECT(\"1:5\"))-1)"
          f"/(1+Ke)^ROW(INDIRECT(\"1:5\")))+({fin['ROEs']}-Ke)*{fin['BVPS']}*(1+{fin['ROEs']}*(1-{fin['Payout']}))^4*(1+{fin['FINg']})"
          f"/(Ke-{fin['FINg']})/(1+Ke)^5,\"\")")
    C.put(ws, r, 3, ri, C.NUM2)
    out["ri"] = f"Valuation!$C${r}"
    r += 2

    # Multiples & normalized earnings (peer medians from Peers sheet)
    r = C.section(ws, r, "Relative value cross-checks (peer medians from the Peers sheet)")
    annual_letters = letters[:-1] if d["fundamentals"]["columns"][-1] == "TTM" else letters
    om = metric_rows["operating_margin"]
    rows_x = [
        ("Fwd P/E: peer median x forward EPS", "=IFERROR(PeerMed_fwd_pe*FwdEPS,\"\")"),
        ("EV/EBITDA: peer median x EBITDA - net debt", f"=IFERROR((PeerMed_ev_ebitda*Ratios!{lvl}{metric_rows['ebitda']}-NetDebt-Minority)/Shares,\"\")"),
        ("EV/Sales: peer median x revenue - net debt", "=IFERROR((PeerMed_ev_sales*RevBase-NetDebt-Minority)/Shares,\"\")"),
        ("Own-history median P/E x latest GAAP EPS (re-rating to its own past)", f"=IFERROR(OwnMedPE*Financials!{lvl}{fin_rows['eps_diluted']},\"\")"),
        ("Normalized EPS (avg EBIT margin x revenue, less interest, after tax)",
         f"=IFERROR((RevBase*AVERAGE(Ratios!{annual_letters[0]}{om}:{annual_letters[-1]}{om})-N(Financials!{lvl}{fin_rows['interest_expense']}))*(1-TaxRate)/Shares,\"\")"),
    ]
    for lab, f in rows_x:
        C.put(ws, r, 2, lab)
        C.put(ws, r, 3, f, C.NUM2)
        r += 1
    C.put(ws, r, 2, "Normalized EPS x peer median P/E (cyclicals)", bold=True)
    C.put(ws, r, 3, f"=IFERROR(C{r - 1}*PeerMed_fwd_pe,\"\")", C.NUM2, bold=True)
    out["norm"] = f"Valuation!$C${r}"
    out["mult_pe"] = f"Valuation!$C${r - 5}"
    out["mult_evs"] = f"Valuation!$C${r - 3}"
    r += 2

    # life-cycle method map
    stage = (notes.get("lifecycle") or {}).get("stage") or d["lifecycle"]["suggested"]
    r = C.section(ws, r, f"Method by life-cycle stage (current: {stage})")
    C.header(ws, r, ["", "Method", "Value / share", "Primary for this stage?"], 1)
    r += 1
    methods = [
        ("3-stage DCF (base)", out["vps"]["Base"], ("High growth", "Mature")),
        ("DCF probability-weighted", "PWValue", ("High growth", "Mature", "Cyclical")),
        ("DDM (2-stage)", out["ddm"], ("Mature dividend payer",)),
        ("Justified P/B", out["jpb"], ("Financials",)),
        ("Residual income", out["ri"], ("Financials",)),
        ("Peer fwd P/E", out["mult_pe"], ("Mature",)),
        ("Peer EV/Sales", out["mult_evs"], ("Early / unprofitable",)),
        ("Normalized earnings x P/E", out["norm"], ("Cyclical",)),
    ]
    for lab, cell, stages in methods:
        C.put(ws, r, 2, lab)
        C.put(ws, r, 3, f"={cell}", C.NUM2)
        C.put(ws, r, 4, "Yes" if stage in stages else "")
        r += 1
    C.status_colors(ws, f"D{r - len(methods)}:D{r}")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 3, 58
    for col in range(3, 15):
        ws.column_dimensions[L(col)].width = 13
    return out


def sheet_reverse(wb, d, refs) -> None:
    ws = wb.create_sheet("Reverse DCF")
    sc = refs["scen"]
    r = C.title(ws, "Reverse DCF: what growth is priced in?",
                "Each row runs the base model with a different stage-1 revenue growth. The row closest to today's "
                "price (highlighted) is the growth the market is paying for.")
    rev = d.get("reverse_dcf") or {}
    C.put(ws, r, 2, "Implied stage-1 growth (Python solve, base inputs at build time)")
    C.put(ws, r, 3, rev.get("implied_g1"), C.PCT, bold=True)
    C.put(ws, r, 4, rev.get("note", "")).font = C.MUTED
    r += 1
    C.put(ws, r, 2, "Compare with consensus 2-yr EPS growth and the company's own 3-yr revenue CAGR.").font = C.MUTED
    r += 2
    hdr = r
    labels = ["Stage-1 growth", *[f"Rev {t}" for t in range(1, 11)], *[f"FCFF {t}" for t in range(1, 11)],
              "Terminal value", "EV", "Value / share", "Gap vs price"]
    C.header(ws, hdr, labels, 1)
    g = lambda k: sc[("Base", k)]  # noqa: E731
    growths = [x / 1000 for x in range(-100, 401, 25)]
    for i, gr in enumerate(growths):
        rr = hdr + 1 + i
        C.put(ws, rr, 1, gr, C.PCT, inp=True)
        for t in range(1, 11):
            col = 1 + t
            prev = "RevBase" if t == 1 else f"{L(col - 1)}{rr}"
            gt = f"IF({t}<={g('n1')},$A{rr},$A{rr}-($A{rr}-{g('g_terminal')})*({t}-{g('n1')})/(10-{g('n1')}))"
            C.put(ws, rr, col, f"={prev}*(1+{gt})", C.MILL)
            fcol = 11 + t
            mt = f"(MarginBase+({g('margin_target')}-MarginBase)*MIN({t},{g('margin_years')})/{g('margin_years')})"
            prev_rev = "RevBase" if t == 1 else f"{L(col - 1)}{rr}"
            C.put(ws, rr, fcol, f"={L(col)}{rr}*{mt}*(1-TaxRate)-({L(col)}{rr}-{prev_rev})/{g('sales_to_capital')}", C.MILL)
        C.put(ws, rr, 22, f"=K{rr}*(1+{g('g_terminal')})*{g('margin_target')}*(1-TaxRate)*(1-{g('g_terminal')}/{g('ronic')})"
                          f"/(WACC+{g('wacc_adj')}-{g('g_terminal')})", C.MILL)
        C.put(ws, rr, 23, f"=NPV(WACC+{g('wacc_adj')},L{rr}:U{rr})+V{rr}/(1+WACC+{g('wacc_adj')})^10", C.MILL)
        C.put(ws, rr, 24, f"=(W{rr}-NetDebt-Minority)/Shares", C.NUM2)
        C.put(ws, rr, 25, f"=X{rr}/Price-1", C.PCT)
    first, last = hdr + 1, hdr + len(growths)
    C.highlight_formula(ws, f"A{first}:Y{last}", f"ABS($X{first}-Price)=MIN(INDEX(ABS($X${first}:$X${last}-Price),0))")
    ws.column_dimensions["A"].width = 14
    for col in range(2, 26):
        ws.column_dimensions[L(col)].width = 11
    C.freeze(ws, f"B{first}")


# ---------------------------------------------------------------------------
# Peers
# ---------------------------------------------------------------------------

PEER_COLS = [("ticker", "Ticker", None), ("name", "Name", None), ("price", "Price", C.NUM2),
             ("market_cap", "Market cap", C.MILL), ("fwd_pe", "Fwd P/E", '0.0"x"'), ("ttm_pe", "TTM P/E (GAAP)", '0.0"x"'), ("peg", "PEG (fwd P/E / 2y EPS CAGR)", "0.00"),
             ("ev_ebitda", "EV/EBITDA", '0.0"x"'), ("ev_ebit", "EV/EBIT", '0.0"x"'), ("ev_sales", "EV/Sales", '0.00"x"'),
             ("fcf_yield", "FCF yield", C.PCT), ("p_b", "P/B", '0.00"x"'), ("p_tbv", "P/TBV", '0.00"x"'),
             ("p_ffo", "P/FFO (approx)", '0.0"x"'), ("dividend_yield", "Div. yield", C.PCT), ("roic", "ROIC", C.PCT),
             ("roe", "ROE", C.PCT), ("operating_margin", "EBIT margin", C.PCT), ("revenue_growth", "Revenue growth", C.PCT),
             ("eps_growth_2y_cagr", "EPS growth 2y CAGR (cons.)", C.PCT), ("net_debt_to_ebitda", "Net debt/EBITDA", '0.0"x"'),
             ("eps_rev_90d", "EPS rev. 90d", C.PCT), ("analysts", "Analysts", "0"), ("flags", "Flags", None)]
MEDIAN_NAMES = {"fwd_pe", "ev_ebitda", "ev_sales", "peg", "ev_ebit", "p_b", "p_tbv", "p_ffo", "fcf_yield", "roic"}


def sheet_peers(wb, d, metric_rows, letters) -> int:
    ws = wb.create_sheet("Peers")
    r = C.title(ws, "Relative value vs the full sub-industry peer group",
                "Compare with the MEDIAN of real peers (not the average of your own picks). Peer data uses Yahoo "
                "definitions; the target's own figures on other sheets come from SEC filings.")
    peers = d.get("peers") or []
    C.header(ws, r, [lab for _, lab, _ in PEER_COLS], 1)
    first = r + 1
    for i, p in enumerate(peers):
        rr = first + i
        for j, (k, _, fmt) in enumerate(PEER_COLS):
            c = C.put(ws, rr, 1 + j, p.get(k), fmt)
            if i == 0:
                c.font = C.BOLD
    last = first + len(peers) - 1
    med_row = last + 1
    prem_row = last + 2
    C.put(ws, med_row, 1, "Peer median (excl. target)", bold=True)
    C.put(ws, prem_row, 1, f"{d['ticker']} premium / (discount) vs median", bold=True)
    for j, (k, _, fmt) in enumerate(PEER_COLS):
        if j < 2 or k in ("flags", "ticker", "name"):
            continue
        no_premium = k in ("price", "market_cap", "analysts")
        le = L(1 + j)
        if len(peers) > 1:
            C.put(ws, med_row, 1 + j, f"=IFERROR(MEDIAN({le}{first + 1}:{le}{last}),\"\")", fmt, bold=True)
            if not no_premium:
                C.put(ws, prem_row, 1 + j, f"=IFERROR({le}{first}/{le}{med_row}-1,\"\")", C.PCT)
        if k in MEDIAN_NAMES:
            C.name(wb, f"PeerMed_{k}", "Peers", med_row, 1 + j)
    if len(peers) <= 1:
        C.put(ws, med_row, 3, "No peers: run `iat company <T> --peers A,B,C` or set the sub-industry").font = C.MUTED
    for k in MEDIAN_NAMES:
        if f"PeerMed_{k}" not in wb.defined_names:
            C.put(ws, med_row + 5, 30, None)
            C.name(wb, f"PeerMed_{k}", "Peers", med_row + 5, 30)

    r = prem_row + 3
    r = C.section(ws, r, "Own valuation history (premium or discount to itself)")
    annual = [c for c in d["ratios"]["columns"] if c != "TTM"]
    C.header(ws, r, ["Metric", *annual, "Median", "Current", "Current vs median"], 1)
    pe_row = metric_rows["pe_at_period_end"]
    fy_row = metric_rows["fcf_yield_at_period_end"]
    for i, (lab, mrow, fmt) in enumerate([("P/E at fiscal year end", pe_row, '0.0"x"'),
                                          ("FCF yield at fiscal year end", fy_row, C.PCT)]):
        rr = r + 1 + i
        C.put(ws, rr, 1, lab)
        for j in range(len(annual)):
            C.put(ws, rr, 2 + j, f"=Ratios!{letters[j]}{mrow}", fmt)
        mcol = 2 + len(annual)
        C.put(ws, rr, mcol, f"=IFERROR(MEDIAN(B{rr}:{L(mcol - 1)}{rr}),\"\")", fmt, bold=True)
        if i == 0:
            C.name(wb, "OwnMedPE", "Peers", rr, mcol)
            C.put(ws, rr, mcol + 1, f"=IFERROR({L(1 + [k for k, _, _ in PEER_COLS].index('ttm_pe'))}{first}+0,\"\")",
                  fmt, comment="Current trailing GAAP P/E (Yahoo), comparable with the year-end history")
        else:
            C.put(ws, rr, mcol + 1, f"=IFERROR({L(1 + [k for k, _, _ in PEER_COLS].index('fcf_yield'))}{first},\"\")", fmt)
        C.put(ws, rr, mcol + 2, f"=IFERROR({L(mcol + 1)}{rr}/{L(mcol)}{rr}-1,\"\")", C.PCT)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 30
    for col in range(3, len(PEER_COLS) + 2):
        ws.column_dimensions[L(col)].width = 12
    ws.column_dimensions[L(len(PEER_COLS))].width = 50
    C.freeze(ws, f"C{first}")
    return prem_row


# ---------------------------------------------------------------------------
# Gate
# ---------------------------------------------------------------------------

def sheet_gate(wb, d, notes, metric_rows, letters, cols) -> str:
    ws = wb.create_sheet("Gate")
    th = wb._iat_settings["thresholds"]
    r = C.title(ws, f"Earnings quality and credit gate: {d['ticker']}",
                "Stop here on FAIL. WATCH items need a written reason before moving on. 'TO REVIEW' items come "
                "from reading the 10-K: the AI drafts, you confirm.")
    C.header(ws, r, ["Test", "Value", "Pass threshold", "Fail threshold", "Status", "Rule", "Evidence / comment (cite the filing)"], 1)
    first = r + 1
    rr = first
    annual = [letters[i] for i, c in enumerate(cols) if c != "TTM"]
    fin = d["profile"] in ("bank", "insurer")
    overrides = {o.get("test"): o for o in (notes.get("gate_review") or [])}
    for spec in GATE_SPEC:
        if fin and spec.get("skip_fin"):
            continue
        mr = metric_rows[spec["ratio"]]
        if spec["agg"] == "avg3":
            rng = ",".join(f"Ratios!{le}{mr}" for le in annual[-3:])
            val = f"=IFERROR(AVERAGE({rng}),\"\")"
        elif spec["agg"] == "level":
            val = f"=IF(ISNUMBER(Ratios!{letters[-1]}{mr}),Ratios!{letters[-1]}{mr},IF(ISNUMBER(Ratios!{annual[-1]}{mr}),Ratios!{annual[-1]}{mr},\"\"))"
        else:
            val = f"=Ratios!{annual[-1]}{mr}"
        if spec["kind"] == "roic_wacc":
            val = f"=IFERROR(Ratios!{annual[-1]}{mr}-WACC,\"\")"
        p, f = thr(th, spec.get("p")), thr(th, spec.get("f"))
        if spec["kind"] == "roic_wacc":
            p = 0.0
        fmt = ratios.RATIOS[spec["ratio"]].fmt if spec["kind"] != "roic_wacc" else C.PCT
        C.put(ws, rr, 1, spec["test"])
        C.put(ws, rr, 2, val, fmt)
        C.put(ws, rr, 3, p, fmt, inp=True)
        C.put(ws, rr, 4, f, fmt, inp=f is not None)
        k = spec["kind"]
        B, Cc, D = f"B{rr}", f"C{rr}", f"D{rr}"
        if k == "hi":
            st = f"=IF(NOT(ISNUMBER({B})),\"NO DATA\",IF({B}>={Cc},\"PASS\",IF({B}<{D},\"FAIL\",\"WATCH\")))"
        elif k == "lo":
            st = f"=IF(NOT(ISNUMBER({B})),\"NO DATA\",IF({B}<={Cc},\"PASS\",IF({B}>{D},\"FAIL\",\"WATCH\")))"
        elif k == "watch_above":
            st = f"=IF(NOT(ISNUMBER({B})),\"NO DATA\",IF({B}>{Cc},\"WATCH\",\"PASS\"))"
        else:  # pass_at_least, roic_wacc
            st = f"=IF(NOT(ISNUMBER({B})),\"NO DATA\",IF({B}>={Cc},\"PASS\",\"WATCH\"))" if k == "pass_at_least" else \
                 f"=IF(NOT(ISNUMBER({B})),\"NO DATA\",IF({B}>{Cc},\"PASS\",\"WATCH\"))"
        C.put(ws, rr, 5, st)
        from iat.company import rule_text
        C.put(ws, rr, 6, rule_text(spec, th)).font = C.MUTED
        o = overrides.get(spec["test"], {})
        C.put(ws, rr, 7, o.get("evidence") or spec.get("note", ""), wrap=True, ai=bool(o.get("evidence")))
        rr += 1
    for g in d["gate"]:
        if g["auto"] and g["test"].startswith("SEC filing red flag"):
            C.put(ws, rr, 1, g["test"])
            C.put(ws, rr, 5, g["status"], inp=True)
            C.put(ws, rr, 6, g["rule"]).font = C.MUTED
            o = overrides.get(g["test"], {})
            C.put(ws, rr, 7, o.get("evidence", "Open the 8-K and explain"), wrap=True, ai=bool(o))
            rr += 1
    for g in d["gate"]:
        if g["auto"]:
            continue
        o = overrides.get(g["test"], {})
        C.put(ws, rr, 1, g["test"])
        cell = C.put(ws, rr, 5, o.get("status", "TO REVIEW"), inp=True)
        C.dropdown(ws, cell.coordinate, ["PASS", "WATCH", "FAIL", "TO REVIEW", "N/A"])
        C.put(ws, rr, 6, g["rule"]).font = C.MUTED
        C.put(ws, rr, 7, o.get("evidence", ""), wrap=True, ai=bool(o.get("evidence")))
        rr += 1
    last = rr - 1
    rr += 1
    C.put(ws, rr, 1, "OVERALL GATE", bold=True)
    rng = f"E{first}:E{last}"
    C.put(ws, rr, 5, f"=IF(COUNTIF({rng},\"FAIL\")>0,\"FAIL\",IF(COUNTIF({rng},\"TO REVIEW\")>0,\"INCOMPLETE\","
                     f"IF(COUNTIF({rng},\"WATCH\")>0,\"WATCH\",\"PASS\")))", bold=True)
    C.name(wb, "GateStatus", "Gate", rr, 5)
    C.put(ws, rr + 1, 1, "Counts: PASS / WATCH / FAIL / TO REVIEW / NO DATA")
    C.put(ws, rr + 1, 5, f"=COUNTIF({rng},\"PASS\")&\" / \"&COUNTIF({rng},\"WATCH\")&\" / \"&COUNTIF({rng},\"FAIL\")"
                         f"&\" / \"&COUNTIF({rng},\"TO REVIEW\")&\" / \"&COUNTIF({rng},\"NO DATA\")")
    C.status_colors(ws, f"E{first}:E{rr}")
    for col, w in zip("ABCDEFG", (48, 12, 14, 14, 14, 34, 80)):
        ws.column_dimensions[col].width = w
    C.freeze(ws, f"B{first}")
    return "GateStatus"


# ---------------------------------------------------------------------------
# AI-authored sheets
# ---------------------------------------------------------------------------

def _rows_sheet(wb, name, heading, sub, columns, rows, widths, pending_hint):
    ws = wb.create_sheet(name)
    r = C.title(ws, heading, sub)
    C.header(ws, r, [c[1] for c in columns], 1)
    if not rows:
        C.put(ws, r + 1, 1, pending_hint, ai=True)
    for i, item in enumerate(rows, start=1):
        for j, (k, _) in enumerate(columns):
            v = item.get(k) if isinstance(item, dict) else (item if j == 0 else "")
            C.put(ws, r + i, 1 + j, v, wrap=True, ai=True)
    for j, w in enumerate(widths):
        ws.column_dimensions[L(1 + j)].width = w
    return ws


def ai_sheets(wb, d, notes) -> None:
    pend = "Pending: run /investment-analyst:pick-the-stock, which writes notes.json, then `iat excel company <T>`"
    b = notes.get("business") or {}
    ws = wb.create_sheet("Business & Moat")
    r = C.title(ws, "Business and moat", "A moat claim needs numbers behind it: ROIC above peers for 5+ years.")
    C.put(ws, r, 1, "What the company does", bold=True)
    C.put(ws, r, 2, b.get("description", pend), wrap=True, ai=True)
    r += 2
    C.header(ws, r, ["Segment", "Revenue share", "Operating profit share", "Trend / comment", "Source"], 1)
    for i, s in enumerate(b.get("segments") or [], start=1):
        for j, k in enumerate(("name", "revenue_share", "profit_share", "comment", "source")):
            C.put(ws, r + i, 1 + j, s.get(k), C.PCT if "share" in k else None, wrap=True, ai=True)
    r += max(len(b.get("segments") or []), 1) + 2
    C.header(ws, r, ["Moat source", "Rating", "Evidence", "Confirmed in numbers?", "Source"], 1)
    for i, mo in enumerate(notes.get("moat") or [], start=1):
        for j, k in enumerate(("factor", "rating", "evidence", "numbers", "source")):
            C.put(ws, r + i, 1 + j, mo.get(k), wrap=True, ai=True)
    if not notes.get("moat"):
        for i, f in enumerate(["Switching costs", "Scale / cost advantage", "Network effects", "Brand",
                               "Regulatory barriers / licenses", "Intangibles / IP"], start=1):
            C.put(ws, r + i, 1, f)
            C.put(ws, r + i, 2, "", ai=True)
    for col, w in zip("ABCDE", (34, 16, 70, 30, 40)):
        ws.column_dimensions[col].width = w

    _rows_sheet(wb, "SWOT", "SWOT tied to model assumptions",
                "Each item must change a number (Inputs sheet) or it does not belong here.",
                [("type", "S/W/O/T"), ("item", "Item"), ("evidence", "Evidence"), ("source", "Source"),
                 ("assumption", "Assumption it drives"), ("direction", "Direction"), ("scenario", "Scenario")],
                notes.get("swot") or [], (8, 40, 60, 34, 26, 10, 12), pend)

    ws = wb.create_sheet("Bear Case")
    r = C.title(ws, "Bear case first, then the pre-mortem",
                "Written before the bull case to counter confirmation bias.")
    for heading, key in (("Bear case", "bear_case"), ("Pre-mortem: it is 6 months from now and the stock fell 25%. What happened?", "premortem"),
                         ("Thesis stops: fundamental changes that would make you wrong", "thesis_stops")):
        C.put(ws, r, 1, heading, bold=True)
        r += 1
        items = notes.get(key) or [pend]
        for it in items:
            C.put(ws, r, 1, it if isinstance(it, str) else it.get("text"), wrap=True, ai=True)
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 140

    fr = notes.get("filing_review") or []
    ws = _rows_sheet(wb, "Filing Review", "10-K / 10-Q / proxy reading log",
                     "Sections: MD&A and critical estimates, segment note, commitments & contingencies, debt note, "
                     "auditor's report and CAMs, Item 9A, revenue recognition, related parties, proxy pay metrics.",
                     [("section", "Section"), ("finding", "Finding"), ("risk", "Risk (Low/Med/High)"),
                      ("source", "Document"), ("quote", "Short quote / page")], fr, (30, 80, 14, 34, 60), pend)
    fp = d.get("filings") or {}
    r = ws.max_row + 3
    if fp.get("eight_k"):
        C.put(ws, r, 1, "8-K filings, last 9 months (auto)", bold=True)
        r += 1
        for e in fp["eight_k"]:
            C.put(ws, r, 1, e["filed"])
            C.put(ws, r, 2, ", ".join(e["descriptions"]), wrap=True)
            C.put(ws, r, 4, e["url"])
            r += 1

    _rows_sheet(wb, "Catalysts", "Catalysts inside the horizon",
                "Without a dated catalyst, a cheap stock can stay cheap past your 6-month window.",
                [("event", "Event"), ("date", "Date / window"), ("within_window", "Inside horizon?"),
                 ("impact", "Expected impact"), ("source", "Source")], notes.get("catalysts") or [],
                (46, 18, 14, 60, 50), pend)


# ---------------------------------------------------------------------------
# Decision, Summary, Checks, Sources
# ---------------------------------------------------------------------------

def sheet_decision(wb, d, notes, val) -> None:
    ws = wb.create_sheet("Decision")
    v = wb._iat_settings["valuation"]
    cats = notes.get("catalysts") or []
    has_cat = any(str(c.get("within_window", "")).lower() in ("yes", "true", "y") for c in cats)
    r = C.title(ws, "Decision: expected 6-month return, catalyst, size, stops",
                "Intrinsic value gives the margin of safety; the catalyst is what closes the gap inside 6 months.")
    C.put(ws, r, 2, "Gap closure inside the horizon")
    C.put(ws, r, 3, v["gap_closure_with_catalyst"] if has_cat else v["gap_closure_without_catalyst"], "0%", inp=True)
    C.put(ws, r, 4, "Share of (value - price) you expect to close in 6 months. Default depends on a dated catalyst.").font = C.MUTED
    closure = f"$C${r}"
    r += 2
    C.header(ws, r, ["", "Scenario", "Probability", "Intrinsic value", "6-month target", "Return"], 1)
    first = r + 1
    for i, s in enumerate(("Bear", "Base", "Bull")):
        rr = first + i
        C.put(ws, rr, 2, s)
        C.put(ws, rr, 3, f"={wb._iat_refs['scen'][(s, 'probability')]}", "0%")
        C.put(ws, rr, 4, f"={val['vps'][s]}", C.NUM2)
        C.put(ws, rr, 5, f"=Price+(D{rr}-Price)*{closure}", C.NUM2)
        C.put(ws, rr, 6, f"=E{rr}/Price-1", C.PCT)
    rr = first + 3
    C.put(ws, rr, 2, "Expected 6-month return", bold=True)
    C.put(ws, rr, 6, f"=SUMPRODUCT(C{first}:C{first + 2},F{first}:F{first + 2})", C.PCT, bold=True)
    C.name(wb, "ExpReturn", "Decision", rr, 6)
    C.put(ws, rr + 1, 2, "After short-term capital gains tax")
    C.put(ws, rr + 1, 6, "=ExpReturn*(1-STTax)", C.PCT)
    C.put(ws, rr + 2, 2, "Expected 6-month price")
    C.put(ws, rr + 2, 6, "=Price*(1+ExpReturn)", C.NUM2)
    r = rr + 4
    r = C.section(ws, r, "Price stop and position size (volatility-based)")
    size_start = r
    lines = [
        ("Entry reference price", "=Price", C.NUM2),
        ("Price stop = entry - k x ATR", "=Price-ATRk*ATR", C.NUM2),
        ("Risk per share", f"=C{r}-C{r + 1}", C.NUM2),
        ("Shares by risk budget", f"=IF(N(Portfolio)=0,\"set Portfolio on Inputs\",ROUNDDOWN(Portfolio*RiskPct/C{r + 2},0))", C.NUM0),
        ("Shares by max-position cap", "=IF(N(Portfolio)=0,\"set Portfolio on Inputs\",ROUNDDOWN(Portfolio*MaxPct/Price,0))", C.NUM0),
        ("Shares (smaller of the two)", f"=IFERROR(MIN(C{r + 3},C{r + 4}),\"\")", C.NUM0),
        ("Position value", f"=IFERROR(C{r + 5}*Price,\"\")", C.NUM0),
        ("% of portfolio", f"=IFERROR(C{r + 6}/Portfolio,\"\")", C.PCT),
        ("Loss if stopped (% of portfolio)", f"=IFERROR(C{r + 5}*C{r + 2}/Portfolio,\"\")", "0.00%"),
        ("Reward / risk (expected gain / stop distance)", f"=IFERROR((Price*ExpReturn)/C{r + 2},\"\")", "0.0"),
    ]
    for i, (lab, f, fmt) in enumerate(lines):
        C.put(ws, r + i, 2, lab)
        C.put(ws, r + i, 3, f, fmt)
    r += len(lines) + 1
    r = C.section(ws, r, "Scale-in plan (spreads entry timing risk)")
    C.header(ws, r, ["", "Tranche", "Trigger", "Shares", "Reference price"], 1)
    sh = f"$C${size_start + 5}"
    tr = [("1/3", "Thesis in place (this workbook + macro scorecard >= 3 of 4)", f"=IFERROR(ROUNDDOWN({sh}/3,0),\"\")", "=Price"),
          ("2/3", "Relative strength confirms (sector ETF/SPY above 10w and 40w MAs)", f"=IFERROR(ROUNDDOWN({sh}/3,0),\"\")", "market"),
          ("3/3", "Pullback toward entry - 1 x ATR", f"=IFERROR({sh}-2*ROUNDDOWN({sh}/3,0),\"\")", "=Price-ATR")]
    for i, (t, trig, shares, px) in enumerate(tr, start=1):
        C.put(ws, r + i, 2, t)
        C.put(ws, r + i, 3, trig, wrap=True)
        C.put(ws, r + i, 4, shares, C.NUM0)
        C.put(ws, r + i, 5, px, C.NUM2)
    r += 5
    r = C.section(ws, r, "Exit rules written before buying")
    C.put(ws, r, 2, "Thesis stop(s)")
    C.put(ws, r, 3, "\n".join(x if isinstance(x, str) else x.get("text", "") for x in (notes.get("thesis_stops") or [])) or "Write them before buying",
          wrap=True, inp=True)
    C.put(ws, r + 1, 2, "Price stop")
    C.put(ws, r + 1, 3, f"=C{size_start + 1}", C.NUM2)
    C.put(ws, r + 2, 2, "Catalyst and date")
    C.put(ws, r + 2, 3, "; ".join(f"{c.get('event')} ({c.get('date')})" for c in cats) or "None identified", wrap=True, ai=bool(cats))
    r += 4
    r = C.section(ws, r, "Your decision")
    C.put(ws, r, 2, "Gate status")
    C.put(ws, r, 3, "=GateStatus", bold=True)
    C.put(ws, r + 1, 2, "Decision")
    cell = C.put(ws, r + 1, 3, "", inp=True)
    C.dropdown(ws, cell.coordinate, ["BUY (tranche 1)", "WATCH", "PASS"])
    C.put(ws, r + 2, 2, "Why (your words)")
    C.put(ws, r + 2, 3, "", inp=True, wrap=True)
    C.put(ws, r + 3, 2, "Date")
    C.put(ws, r + 3, 3, "", inp=True)
    C.status_colors(ws, f"C{r}:C{r}")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 3, 46
    ws.column_dimensions["C"].width, ws.column_dimensions["D"].width = 60, 14
    ws.column_dimensions["E"].width = ws.column_dimensions["F"].width = 16
    ws.row_dimensions[r + 2].height = 60


def sheet_summary(wb, d, notes, val) -> None:
    ws = wb.create_sheet("Summary")
    stage = (notes.get("lifecycle") or {}).get("stage") or d["lifecycle"]["suggested"]
    r = C.title(ws, f"{d['ticker']} - {d.get('name')}",
                f"{d['sector']} / {d.get('sub_industry')} / built {d['as_of'][:10]} / statements: {d['source'].upper()}")
    items = [
        ("Price", "=Price", C.NUM2),
        ("Gate (earnings quality + credit)", "=GateStatus", None),
        ("Life-cycle stage", stage, None),
        ("Base DCF value / share", f"={val['vps']['Base']}", C.NUM2),
        ("Probability-weighted value / share", "=PWValue", C.NUM2),
        ("Upside to probability-weighted value", "=PWValue/Price-1", C.PCT),
        ("Reverse DCF: implied stage-1 growth", d.get("reverse_dcf", {}).get("implied_g1"), C.PCT),
        ("Fwd P/E vs peer median", f"=IFERROR(Peers!{L(1 + [k for k, _, _ in PEER_COLS].index('fwd_pe'))}{wb._iat_prem_row},\"\")", C.PCT),
        ("Expected 6-month return", "=ExpReturn", C.PCT),
        ("WACC", "=WACC", "0.00%"),
        ("Rate sensitivity (% per +100bp in 10Y)", (d.get("rate_sensitivity") or {}).get("pct_move_per_100bp"), "0.0"),
        ("Next earnings date", d["market"].get("next_earnings"), None),
    ]
    for i, (lab, v, fmt) in enumerate(items):
        C.put(ws, r + i, 1, lab, bold=True)
        C.put(ws, r + i, 2, v, fmt)
    C.status_colors(ws, f"B{r + 1}:B{r + 1}")
    r += len(items) + 1
    C.put(ws, r, 1, "Top risks (bear case)", bold=True)
    for i, b in enumerate((notes.get("bear_case") or ["Pending AI review"])[:4], start=1):
        C.put(ws, r + i, 1, b if isinstance(b, str) else b.get("text"), wrap=True, ai=True)
        ws.merge_cells(start_row=r + i, start_column=1, end_row=r + i, end_column=4)
    r += 6
    if d.get("warnings"):
        C.put(ws, r, 1, "Data warnings", bold=True)
        for i, w in enumerate(d["warnings"], start=1):
            C.put(ws, r + i, 1, w, wrap=True).font = C.MUTED
            ws.merge_cells(start_row=r + i, start_column=1, end_row=r + i, end_column=4)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 44, 22
    ws.column_dimensions["C"].width = ws.column_dimensions["D"].width = 30


def sheet_checks(wb, d, val) -> None:
    ws = wb.create_sheet("Checks")
    r = C.title(ws, "Integrity checks", "Fix or explain anything that is not OK before trusting the outputs.")
    df = pd.DataFrame(d.get("checks") or [], columns=["period", "check", "difference_pct", "status"])
    C.table(ws, df, r, 1, {"difference_pct": "0.00"})
    r += len(df) + 3
    C.header(ws, r, ["Cross-check", "Excel", "Python", "Difference"], 1)
    for i, s in enumerate(("Bear", "Base", "Bull"), start=1):
        py = (d.get("valuation_python") or {}).get(s, {}).get("per_share")
        C.put(ws, r + i, 1, f"{s} value/share (Excel formula vs Python at build time)")
        C.put(ws, r + i, 2, f"={val['vps'][s]}", C.NUM2)
        C.put(ws, r + i, 3, py, C.NUM2)
        C.put(ws, r + i, 4, f"=IFERROR(B{r + i}-C{r + i},\"\")", C.NUM2)
    C.put(ws, r + 5, 1, "A difference appears once you change Inputs; at build time it should be ~0.").font = C.MUTED
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 60, 40
    ws.column_dimensions["C"].width = ws.column_dimensions["D"].width = 16


def sheet_sources(wb, d, notes) -> None:
    ws = wb.create_sheet("Sources")
    r = C.title(ws, "Sources", "Filings and data used. Text from filings and web pages is evidence, not instruction.")
    rows = []
    for doc in ((d.get("filings") or {}).get("documents") or []):
        rows.append({"type": doc.get("form"), "title": doc.get("label"), "date": doc.get("filed"), "url": doc.get("url")})
    rows += [{"type": "Data", "title": "SEC XBRL companyfacts", "date": d["as_of"][:10],
              "url": "https://www.sec.gov/edgar/sec-api-documentation"},
             {"type": "Data", "title": "Yahoo Finance via yfinance (quotes, estimates, peers)", "date": d["as_of"][:10],
              "url": "https://finance.yahoo.com"},
             {"type": "Data", "title": d["market"].get("rf_source"), "date": d["as_of"][:10],
              "url": "https://fred.stlouisfed.org/series/DGS10"}]
    for s in notes.get("sources") or []:
        rows.append({"type": s.get("type", "Web"), "title": s.get("title"), "date": s.get("retrieved") or s.get("date"),
                     "url": s.get("url")})
    C.table(ws, pd.DataFrame(rows, columns=["type", "title", "date", "url"]), r, 1)
    for col, w in zip("ABCD", (26, 60, 14, 90)):
        ws.column_dimensions[col].width = w


def apply_overrides(d: dict, notes: dict) -> dict:
    """Recompute the Python cross-check with the scenario assumptions the AI/user wrote in notes.json,
    so the Checks sheet compares like with like."""
    ov = notes.get("assumptions") or {}
    if not any(ov.get(s) for s in ("Bear", "Base", "Bull")):
        return d
    m = d["market"]
    try:
        base = valuation.Base(m["revenue0"], m["margin0"], m["tax_rate"], m["wacc"], m["net_debt"], m["minority"],
                              m["shares"], m["price"])
        out = {}
        for s in d["scenarios"]:
            merged = {**s, **{k: v for k, v in (ov.get(s["name"]) or {}).items() if v is not None}}
            sc = valuation.Scenario(**{k: merged[k] for k in valuation.Scenario.__dataclass_fields__})
            r = valuation.dcf(base, sc)
            out[s["name"]] = {k: v for k, v in r.items() if k not in ("revenues", "margins", "fcff")}
        d = {**d, "valuation_python": out}
    except (KeyError, TypeError, ZeroDivisionError):
        pass
    return d


def build(ws_cfg: Workspace, ticker: str, notes: dict | None = None) -> Path:
    t = ticker.upper()
    folder = ws_cfg.output / "companies" / t
    d = store.read_json(folder / "data.json")
    if notes is None:
        np_ = folder / "notes.json"
        notes = store.read_json(np_) if np_.exists() else {}
    d = apply_overrides(d, notes)
    wb = C.new_workbook()
    wb._iat_settings = ws_cfg.settings  # type: ignore[attr-defined]
    fin_rows, cols, letters = sheet_financials(wb, d)
    metric_rows = sheet_ratios(wb, d, fin_rows, cols, letters)
    refs = sheet_inputs(wb, d, notes, fin_rows, metric_rows, cols, letters)
    wb._iat_refs = refs  # type: ignore[attr-defined]
    wb._iat_prem_row = sheet_peers(wb, d, metric_rows, letters)  # type: ignore[attr-defined]
    val = sheet_valuation(wb, d, refs, notes, metric_rows, fin_rows, letters)
    sheet_reverse(wb, d, refs)
    sheet_gate(wb, d, notes, metric_rows, letters, cols)
    ai_sheets(wb, d, notes)
    sheet_decision(wb, d, notes, val)
    sheet_summary(wb, d, notes, val)
    sheet_checks(wb, d, val)
    sheet_sources(wb, d, notes)
    C.readme(wb, f"{t} deep dive", [
        ("Order of work", "1 Business & Moat and SWOT -> 2 Gate (stop on FAIL) -> 3 Peers (relative value) -> "
                          "4 Valuation + Reverse DCF (intrinsic value) -> 5 Decision"),
        ("Inputs", "Change blue cells on Inputs (scenarios, WACC parts) and Decision (gap closure, decision)."),
        ("Green cells", "Written by the AI from filings and news, with sources. Verify before relying on them."),
        ("Financials", "SEC XBRL values; each cell's concept and accession number is on 'XBRL Sources'. "
                       "If a number looks wrong, check it in the 10-K and type the correct value over it."),
        ("Ratios", "Live formulas, so you can audit any cell by clicking it."),
        ("Rebuild", "`iat company <T>` refreshes data; `iat excel company <T>` rebuilds this file with the latest notes.json. "
                    "Rebuilding overwrites manual edits, so save a copy if you typed over values."),
    ])
    order = ["README", "Summary", "Decision", "Gate", "Valuation", "Reverse DCF", "Inputs", "Peers", "Business & Moat",
             "SWOT", "Bear Case", "Filing Review", "Catalysts", "Financials", "Ratios", "Checks", "Sources", "XBRL Sources"]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames] + [s for s in wb._sheets if s.title not in order]
    out = folder / f"{t}_DeepDive_{store.today()}.xlsx"
    wb.save(out)
    return out
