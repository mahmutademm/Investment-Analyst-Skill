"""Screen workbook: Screen_<key>_<date>.xlsx (ranking, metric groups, news check, methodology)."""
from __future__ import annotations

from pathlib import Path

import pandas as pd

from iat import screener, store
from iat.config import Workspace
from iat.excel import common as C

GROUPS = {
    "Valuation": ["ticker", "name", "price", "market_cap", "fwd_pe", "ttm_pe", "own_pe_median_hist", "pe_vs_own_hist",
                  "peg", "eps_growth_2y_cagr", "ev_ebitda", "ev_ebit", "ev_sales", "ev_sales_to_growth", "fcf_yield",
                  "p_b", "p_tbv", "p_ffo", "dividend_yield"],
    "Quality": ["ticker", "roic", "roe", "rotce", "gross_margin", "operating_margin", "op_margin_std", "net_margin_std",
                "cfo_to_ni", "fcf_to_ni", "accruals_ratio", "revenue_growth", "revenue_cagr_3y", "rule_of_40",
                "sbc_pct_revenue", "share_count_change", "capex_to_da"],
    "Expectations": ["ticker", "analysts", "eps_rev_30d", "eps_rev_90d", "eps_rev_30d_fy0", "revisions_up_30d",
                     "revisions_down_30d", "net_revisions_30d", "estimate_dispersion", "beats_4q", "avg_surprise_pct",
                     "next_earnings", "days_to_earnings"],
    "Balance sheet & risk": ["ticker", "net_debt_to_ebitda", "interest_coverage", "debt_to_capital", "current_ratio",
                             "fcf_ttm", "altman_z", "altman_z2", "beneish_m", "piotroski_f", "short_pct_float", "beta",
                             "negative_equity", "country", "currency"],
    "Momentum & trading": ["ticker", "ret_6m", "ret_12m", "rel_6m_vs_etf", "rel_12m_vs_etf", "price_vs_200d", "atr14",
                           "atr14_pct", "adv_usd_20d", "benchmark_etf"],
}
PCT_COLS = {"pe_vs_own_hist", "eps_growth_2y_cagr", "fcf_yield", "dividend_yield", "roic", "roe", "rotce", "gross_margin",
            "operating_margin", "op_margin_std", "net_margin_std", "accruals_ratio", "revenue_growth", "revenue_cagr_3y",
            "sbc_pct_revenue", "share_count_change", "eps_rev_30d", "eps_rev_90d", "eps_rev_30d_fy0", "net_revisions_30d",
            "estimate_dispersion", "debt_to_capital", "short_pct_float", "ret_6m", "ret_12m", "rel_6m_vs_etf",
            "rel_12m_vs_etf", "price_vs_200d", "atr14_pct", "valuation_score", "quality_score", "momentum_score",
            "composite", "data_completeness"}
MILL_COLS = {"market_cap", "fcf_ttm", "adv_usd_20d"}


def _fmts(cols):
    out = {}
    for c in cols:
        if c in PCT_COLS:
            out[c] = C.PCT
        elif c in MILL_COLS:
            out[c] = C.MILL
        elif c in ("ticker", "name", "flags", "country", "currency", "next_earnings", "benchmark_etf"):
            continue
        else:
            out[c] = "0.00"
    return out


def build(ws_cfg: Workspace, key: str, notes: dict | None = None, notes_path: Path | None = None) -> Path:
    df = store.read_latest_csv(ws_cfg, f"screen_{key}")
    if df is None:
        raise RuntimeError(f"No screen data for '{key}'. Run `iat screen {key}` (or wait for the weekly GitHub Action).")
    sub = ws_cfg.sub_industry(key)
    if notes is None and notes_path and notes_path.exists():
        notes = store.read_json(notes_path)
    notes = notes or {}
    as_of = str(df["as_of"].iloc[0])[:16]
    wb = C.new_workbook()

    ws = wb.create_sheet("Ranking")
    r = C.title(ws, f"{sub.get('name', key)}: peer-relative ranking ({len(df)} names, {as_of})",
                "Percentiles are within THIS sub-industry. The rank decides who gets a news check, not what to buy.")
    rank_cols = ["rank", "ticker", "name", "composite", "valuation_score", "quality_score", "momentum_score",
                 "data_completeness", "step1_cheap", "step2_estimates_ok", "step3_temporary_discount",
                 "step4_catalyst_6m", "news_check", "fwd_pe", "eps_rev_30d", "market_cap", "flags"]
    nc = {x.get("ticker"): x for x in notes.get("news_check") or []}
    t = df[[c for c in rank_cols if c in df]].copy()
    if nc:
        t["step3_temporary_discount"] = [nc.get(x, {}).get("classification", "AI news check") for x in t["ticker"]]
        t["step4_catalyst_6m"] = [nc.get(x, {}).get("catalyst_within_window", "AI news check") for x in t["ticker"]]
    first, last = C.table(ws, t, r, 1, _fmts(t.columns), widths={"name": 30, "flags": 70})
    from openpyxl.utils import get_column_letter as L
    for col in ("composite", "valuation_score", "quality_score", "momentum_score"):
        j = list(t.columns).index(col) + 1
        C.heat(ws, f"{L(j)}{first}:{L(j)}{last}")
    for col in ("step1_cheap", "step2_estimates_ok", "step3_temporary_discount", "step4_catalyst_6m"):
        j = list(t.columns).index(col) + 1
        C.status_colors(ws, f"{L(j)}{first}:{L(j)}{last}")
        C.highlight_formula(ws, f"{L(j)}{first}:{L(j)}{last}", f'ISNUMBER(SEARCH("Temporary",{L(j)}{first}))', C.GREEN)
        C.highlight_formula(ws, f"{L(j)}{first}:{L(j)}{last}", f'ISNUMBER(SEARCH("Structural",{L(j)}{first}))', C.RED)
    C.freeze(ws, f"D{first}")

    for name, cols in GROUPS.items():
        s = wb.create_sheet(name)
        cols = [c for c in cols if c in df]
        g = df[cols].copy()
        rr = C.title(s, f"{name} ({key})", "Medians in the last row are the sub-industry yardstick.")
        f0, l0 = C.table(s, g, rr, 1, _fmts(cols), widths={"name": 30})
        C.put(s, l0 + 1, 1, "Median", bold=True)
        for j, c in enumerate(cols):
            if c in ("ticker", "name", "country", "currency", "next_earnings", "benchmark_etf", "negative_equity"):
                continue
            col = L(j + 1)
            C.put(s, l0 + 1, j + 1, f"=IFERROR(MEDIAN({col}{f0}:{col}{l0}),\"\")", _fmts([c]).get(c), bold=True)
        C.freeze(s, f"B{f0}")
        C.autosize(s)

    ws = wb.create_sheet("News Check")
    r = C.title(ws, "Why is it cheap? (top names only)",
                "Sources: last two earnings releases, recent 8-Ks, reputable headlines. Temporary + catalyst inside 6 months = deep-dive candidate.")
    cols = ["ticker", "why_cheap", "classification", "evidence", "catalyst", "catalyst_date", "catalyst_within_window",
            "red_flags", "deep_dive_candidate", "sources"]
    rows = notes.get("news_check") or [{"ticker": x, "classification": "pending"} for x in df[df.get("news_check", "") == "YES"]["ticker"]]
    nd = pd.DataFrame(rows)
    for c in cols:
        if c not in nd:
            nd[c] = ""
    nd["sources"] = nd["sources"].apply(lambda v: "\n".join(v) if isinstance(v, list) else v)
    f0, l0 = C.table(ws, nd[cols], r, 1, wrap_cols=cols[1:], widths={"ticker": 8, "why_cheap": 50, "classification": 14,
                                                                      "evidence": 60, "catalyst": 36, "catalyst_date": 14,
                                                                      "catalyst_within_window": 12, "red_flags": 36,
                                                                      "deep_dive_candidate": 12, "sources": 60})
    for row in ws.iter_rows(min_row=f0, max_row=l0):
        for c in row[1:]:
            c.fill = C.AI_FILL
    C.highlight_formula(ws, f"C{f0}:C{l0}", f'ISNUMBER(SEARCH("Temporary",C{f0}))', C.GREEN)
    C.highlight_formula(ws, f"C{f0}:C{l0}", f'ISNUMBER(SEARCH("Structural",C{f0}))', C.RED)
    r = l0 + 3
    C.put(ws, r, 1, "Your shortlist for the deep dive (3-4 names)", bold=True)
    for i in range(4):
        C.put(ws, r + 1 + i, 1, (notes.get("user_shortlist") or [None] * 4)[i] if i < len(notes.get("user_shortlist") or []) else "", inp=True)

    ws = wb.create_sheet("Methodology")
    r = C.title(ws, "How this screen works", f"Profile: {sub['profile']}. Notes for this sub-industry: {sub.get('notes', '')}")
    spec = screener.PROFILE_METRICS.get(sub["profile"], screener.PROFILE_METRICS["standard"])
    lines = [
        ("Universe", f"Yahoo industries {sub.get('yahoo_industries')} on US exchanges, market cap > "
                     f"${ws_cfg.settings['screening']['min_market_cap'] / 1e9:.0f}B, dollar volume > "
                     f"${ws_cfg.settings['screening']['min_adv_usd'] / 1e6:.0f}M/day. Excluded names are on the Universe sheet."),
        ("Valuation group", ", ".join(f"{k} ({b} is better)" for k, b in spec["valuation"])),
        ("Quality group", ", ".join(f"{k} ({b} is better)" for k, b in spec["quality"])),
        ("Revisions/momentum group", ", ".join(f"{k} ({b} is better)" for k, b in screener.MOMENTUM_METRICS)),
        ("Score", "Percentile within the sub-industry for each metric, averaged per group (needs at least half the metrics), "
                  "composite = average of the 3 groups. Value alone finds traps, momentum alone finds expensive stocks, "
                  "quality alone ignores price."),
        ("4-step logic", "1 cheap (valuation >= median and not above own history) -> 2 estimates not falling (30/90-day FY+1 "
                         "revisions >= -1% and net revisions >= 0) -> 3 discount is temporary (AI news check) -> 4 catalyst inside 6 months."),
        ("Flags", "Flags never remove a name. They tell you what to check: falling estimates, leverage, M-score, accruals, "
                  "downtrend, earnings soon, short interest, foreign filer, partial data."),
        ("Data caveats", "Yahoo data (unofficial). Definitions differ from filings; the deep dive recomputes from SEC XBRL. "
                         "Own-history P/E uses 4-5 fiscal year-ends of GAAP EPS only."),
    ]
    for i, (a, b) in enumerate(lines):
        C.put(ws, r + i, 1, a, bold=True)
        C.put(ws, r + i, 2, b, wrap=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 120

    uni_path = ws_cfg.universes / f"{key}.csv"
    if uni_path.exists():
        ws = wb.create_sheet("Universe")
        u = pd.read_csv(uni_path)
        rr = C.title(ws, "Universe and exclusions", "Edit include/exclude in config/sub_industries.yaml, then rerun.")
        C.table(ws, u, rr, 1, {"market_cap": C.MILL, "adv_usd_3m": C.MILL})
        C.autosize(ws)

    C.readme(wb, f"Sub-industry screen: {sub.get('name', key)}", [
        ("Use", "Ranking -> News Check (AI drafts, you judge) -> pick 3-4 names for /investment-analyst:pick-the-stock."),
        ("Warning", "A screen is cheap and mechanical by design. It narrows the field; it cannot tell you why a stock is cheap."),
    ])
    order = ["README", "Ranking", "News Check", "Valuation", "Quality", "Expectations", "Balance sheet & risk",
             "Momentum & trading", "Methodology", "Universe"]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames]
    out = ws_cfg.output / "screens" / f"Screen_{key}_{store.today()}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
