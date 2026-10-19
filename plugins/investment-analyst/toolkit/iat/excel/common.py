"""Shared styles and helpers for the workbooks."""
from __future__ import annotations

import math
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "1F3864"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
SECTION_FONT = Font(bold=True, size=11, color=NAVY)
INPUT_FONT = Font(color="0000FF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
AI_FILL = PatternFill("solid", fgColor="E2EFDA")
MUTED = Font(color="7F7F7F", italic=True)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
GREEN = PatternFill("solid", fgColor="C6EFCE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")
GREY = PatternFill("solid", fgColor="EDEDED")

PCT, MULT, NUM0, NUM2, MILL = "0.0%", '0.00"x"', "#,##0", "#,##0.00", '#,##0,,"M"'


def clean(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    try:
        import numpy as np
        if isinstance(v, np.generic):
            v = v.item()
            if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
                return None
    except ImportError:
        pass
    return v


def new_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def title(ws, text: str, subtitle: str | None = None, row: int = 1) -> int:
    ws.cell(row=row, column=1, value=text).font = TITLE_FONT
    if subtitle:
        ws.cell(row=row + 1, column=1, value=subtitle).font = MUTED
        return row + 3
    return row + 2


def section(ws, row: int, text: str, col: int = 1) -> int:
    ws.cell(row=row, column=col, value=text).font = SECTION_FONT
    return row + 1


def header(ws, row: int, labels: Iterable[str], col: int = 1) -> None:
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.fill, c.font, c.alignment, c.border = HEADER_FILL, HEADER_FONT, Alignment(wrap_text=True, vertical="center"), BOX


def put(ws, row: int, col: int, value: Any, fmt: str | None = None, *, inp: bool = False, ai: bool = False,
        bold: bool = False, wrap: bool = False, comment: str | None = None):
    c = ws.cell(row=row, column=col, value=clean(value))
    if fmt:
        c.number_format = fmt
    if inp:
        c.font, c.fill = INPUT_FONT, INPUT_FILL
    elif ai:
        c.fill = AI_FILL
    if bold:
        c.font = BOLD
    if wrap:
        c.alignment = WRAP
    if comment:
        c.comment = Comment(comment, "iat")
    c.border = BOX
    return c


def ref(ws_title: str, row: int, col: int, absolute: bool = True) -> str:
    letter = get_column_letter(col)
    sheet = f"'{ws_title}'" if " " in ws_title or "&" in ws_title else ws_title
    return f"{sheet}!${letter}${row}" if absolute else f"{sheet}!{letter}{row}"


def name(wb: Workbook, label: str, sheet: str, row: int, col: int) -> str:
    """Create a workbook-level defined name pointing at one cell; returns the name."""
    target = ref(sheet, row, col)
    dn = DefinedName(label, attr_text=target)
    try:
        wb.defined_names[label] = dn  # openpyxl >= 3.1
    except TypeError:
        wb.defined_names.append(dn)  # older openpyxl
    return label


def table(ws, df: pd.DataFrame, row: int, col: int = 1, formats: dict[str, str] | None = None,
          widths: dict[str, int] | None = None, wrap_cols: Iterable[str] = ()) -> tuple[int, int]:
    """Write a DataFrame with a styled header. Returns (first_data_row, last_data_row)."""
    formats = formats or {}
    wrap_cols = set(wrap_cols)
    header(ws, row, list(df.columns), col)
    for i, (_, r) in enumerate(df.iterrows(), start=1):
        for j, cname in enumerate(df.columns):
            put(ws, row + i, col + j, r[cname], formats.get(cname), wrap=cname in wrap_cols)
    for j, cname in enumerate(df.columns):
        if widths and cname in widths:
            ws.column_dimensions[get_column_letter(col + j)].width = widths[cname]
    return row + 1, row + len(df)


def autosize(ws, min_width: int = 8, max_width: int = 60) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            v = str(c.value)
            if v.startswith("="):
                n = 12
            else:
                n = max(len(x) for x in v.split("\n")) + 2
            widths[c.column] = max(widths.get(c.column, 0), n)
    for col, w in widths.items():
        letter = get_column_letter(col)
        if ws.column_dimensions[letter].width in (None, 13.0) or ws.column_dimensions[letter].width < min_width:
            ws.column_dimensions[letter].width = max(min_width, min(w, max_width))


def status_colors(ws, cell_range: str) -> None:
    ws.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=['"PASS"'], fill=GREEN))
    ws.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=['"WATCH"'], fill=AMBER))
    ws.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=['"FAIL"'], fill=RED))
    ws.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=['"TO REVIEW"'], fill=GREY))
    for word, fill in (("Yes", GREEN), ("No", RED), ("CONFIRMED", GREEN), ("ENTER", GREEN), ("AVOID", RED)):
        ws.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=[f'"{word}"'], fill=fill))


def heat(ws, cell_range: str, reverse: bool = False) -> None:
    lo, hi = ("63BE7B", "F8696B") if reverse else ("F8696B", "63BE7B")
    ws.conditional_formatting.add(cell_range, ColorScaleRule(start_type="min", start_color=lo, mid_type="percentile",
                                                             mid_value=50, mid_color="FFEB84", end_type="max",
                                                             end_color=hi))


def highlight_formula(ws, cell_range: str, formula: str, fill=GREEN) -> None:
    ws.conditional_formatting.add(cell_range, FormulaRule(formula=[formula], fill=fill))


def dropdown(ws, cell: str, options: list[str]) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cell)


def readme(wb: Workbook, heading: str, lines: list[str | tuple[str, str]]) -> None:
    ws = wb.create_sheet("README", 0)
    r = title(ws, heading, "Generated by iat-toolkit. Not investment advice: every decision is yours.")
    for item in lines:
        if isinstance(item, tuple):
            ws.cell(row=r, column=1, value=item[0]).font = BOLD
            ws.cell(row=r, column=2, value=item[1]).alignment = WRAP
        else:
            ws.cell(row=r, column=1, value=item).alignment = WRAP
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Color legend").font = SECTION_FONT
    put(ws, r + 1, 1, "Input you can change", inp=True)
    put(ws, r + 2, 1, "Filled by the AI from filings/news (verify)", ai=True)
    put(ws, r + 3, 1, "=\"Formula (do not overwrite)\"")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110


def freeze(ws, cell: str) -> None:
    ws.freeze_panes = cell
