"""Macro & rotation workbook: Macro_<date>.xlsx (dashboard, cycle rules, rotation, 4-pillar scorecard)."""
from __future__ import annotations

from pathlib import Path

import pandas as pd

from iat import macro, store
from iat.config import SECTOR_ETFS, Workspace
from iat.excel import common as C


def _yes_no(v) -> str:
    s = str(v).strip().lower()
    if s in ("yes", "y", "true", "up", "positive", "supportive", "cheap", "1"):
        return "Yes"
    if s in ("no", "n", "false", "down", "negative", "expensive", "0"):
        return "No"
    return "Neutral" if s in ("flat", "neutral", "mixed") else ""


def build(ws_cfg: Workspace, notes: dict | None = None, notes_path: Path | None = None) -> Path:
    m = store.read_latest_json(ws_cfg, "macro")
    rot = store.read_latest_csv(ws_cfg, "rotation")
    corr = store.read_latest_csv(ws_cfg, "rotation_corr")
    if m is None or rot is None:
        raise RuntimeError("Run `iat macro` and `iat rotation` first (or pull the latest GitHub Actions data).")
    if notes is None and notes_path and notes_path.exists():
        notes = store.read_json(notes_path)
    notes = notes or {}
    wb = C.new_workbook()

    # Dashboard
    ws = wb.create_sheet("Dashboard")
    r = C.title(ws, "Macro dashboard", f"FRED data as of {m['as_of'][:16]}. Changes are in the series' own units.")
    df = pd.DataFrame(m["indicators"])[["group", "label", "latest", "chg_3m", "chg_6m", "chg_12m", "date", "note"]]
    df.columns = ["Group", "Indicator", "Latest", "3m change", "6m change", "12m change", "As of", "How to read it"]
    C.table(ws, df, r, 1, {"Latest": "#,##0.00", "3m change": "+#,##0.00;-#,##0.00", "6m change": "+#,##0.00;-#,##0.00",
                           "12m change": "+#,##0.00;-#,##0.00"},
            widths={"Group": 11, "Indicator": 52, "Latest": 12, "3m change": 11, "6m change": 11, "12m change": 11,
                    "As of": 12, "How to read it": 60})
    C.freeze(ws, f"A{r + 1}")
    if m.get("errors"):
        rr = r + len(df) + 2
        C.put(ws, rr, 1, "Failed series: " + ", ".join(m["errors"])).font = C.MUTED

    # Cycle rules
    ws = wb.create_sheet("Cycle Rules")
    c = m["cycle"]
    final = (notes.get("cycle_call") or {}).get("phase") or ""
    r = C.title(ws, "Business-cycle rule vote (a starting map, not a law)",
                "Each rule votes for a phase when it fires. The suggestion is mechanical; YOUR call goes in the yellow cell.")
    C.put(ws, r, 1, "Rule-based suggestion", bold=True)
    C.put(ws, r, 2, f"{c['suggested_phase']} ({c['confidence_share']:.0%} of votes; runner-up {c['runner_up']})"
          + (" - AMBIGUOUS, read the evidence" if c["ambiguous"] else ""))
    C.put(ws, r + 1, 1, "Your cycle call", bold=True)
    cell = C.put(ws, r + 1, 2, final, inp=True)
    C.dropdown(ws, cell.coordinate, macro.PHASES)
    C.name(wb, "CyclePhase", "Cycle Rules", r + 1, 2)
    C.put(ws, r + 2, 1, "Why", bold=True)
    C.put(ws, r + 2, 2, (notes.get("cycle_call") or {}).get("rationale", ""), wrap=True, ai=True)
    ws.row_dimensions[r + 2].height = 60
    r += 4
    sc = pd.DataFrame([{"Phase": p, "Score": s, "Favored sectors": ", ".join(macro.PHASE_SECTORS[p]),
                        "Usually lag": ", ".join(macro.PHASE_AVOID[p])} for p, s in c["scores"].items()])
    C.table(ws, sc, r, 1)
    r += len(sc) + 2
    rules = pd.DataFrame(m["rules"])[["phase", "weight", "fired", "rule", "evidence"]]
    rules["fired"] = rules["fired"].map({True: "Yes", False: ""})
    rules.columns = ["Votes for", "Weight", "Fired", "Rule", "Evidence"]
    C.table(ws, rules, r, 1)
    C.status_colors(ws, f"C{r + 1}:C{r + len(rules)}")
    for col, w in zip("ABCDE", (22, 70, 10, 70, 50)):
        ws.column_dimensions[col].width = w

    # Rotation
    ws = wb.create_sheet("Rotation")
    r = C.title(ws, "Relative strength vs SPY (entry confirmation pillar)",
                "Market agrees with a thesis when ETF/SPY is above its 10- and 40-week MAs for 3+ weeks, or the RRG "
                "quadrant is Improving/Leading two weeks running. Short spikes (EARLY) are false-start candidates.")
    cols = ["ticker", "name", "level", "status", "rs_pillar", "quadrant", "quadrant_last_week", "quadrant_4w_ago",
            "weeks_above_10w", "above_10w", "above_40w", "rs_ratio", "rs_momentum", "rel_ret_13w_pct", "rel_ret_26w_pct",
            "rel_ret_52w_pct", "rel_to_parent_26w_pct", "price_vs_200d_pct", "atr14_pct", "data_through"]
    t = rot[[x for x in cols if x in rot]].copy()
    t["rs_pillar"] = t["rs_pillar"].map({True: "Yes", False: "No"})
    C.table(ws, t, r, 1, {k: "0.0" for k in ("rs_ratio", "rs_momentum", "rel_ret_13w_pct", "rel_ret_26w_pct",
                                             "rel_ret_52w_pct", "rel_to_parent_26w_pct", "price_vs_200d_pct", "atr14_pct")})
    C.status_colors(ws, f"D{r + 1}:E{r + len(t)}")
    C.heat(ws, f"N{r + 1}:N{r + len(t)}")
    C.autosize(ws)
    C.freeze(ws, f"C{r + 1}")

    # Scorecard
    ws = wb.create_sheet("Scorecard")
    phase = final or c["suggested_phase"]
    r = C.title(ws, "Four-pillar sector scorecard (review monthly, not daily)",
                "Enter a sector only when 3 of 4 pillars agree. Pillars 2 and 4 come from FactSet Earnings Insight / "
                "Yardeni / analyst notes (AI fills with sources); you can overwrite any yellow cell.")
    C.put(ws, r, 1, "Cycle phase used", bold=True)
    C.put(ws, r, 2, "=IF(CyclePhase<>\"\",CyclePhase,\"" + c["suggested_phase"] + " (suggested)\")")
    r += 2
    hdr = ["Sector", "ETF", "1 Cycle fit", "2 Earnings revisions", "3 Relative strength", "4 Valuation vs own history",
           "Pillars agreeing", "Verdict", "Revisions evidence", "Valuation evidence (fwd P/E vs 10y avg; premium to S&P)",
           "Already priced in?", "Proposed weight %", "Names (of target positions)", "Sub-industries the thesis supports"]
    C.header(ws, r, hdr, 1)
    sec_notes = notes.get("sectors") or {}
    alloc = notes.get("allocation") or {}
    subs = notes.get("sub_industries") or {}
    rot_by = {x["ticker"]: x for _, x in rot.iterrows()}
    n_target = ws_cfg.settings["portfolio"]["target_positions"]
    first = r + 1
    for i, (sector, etf) in enumerate(SECTOR_ETFS.items()):
        rr = first + i
        sn = sec_notes.get(sector, {})
        fit = "Yes" if sector in macro.PHASE_SECTORS.get(phase, []) else (
            "No" if sector in macro.PHASE_AVOID.get(phase, []) else "Neutral")
        rv = sn.get("revisions") or {}
        va = sn.get("valuation") or {}
        rs = rot_by.get(etf, {})
        C.put(ws, rr, 1, sector, bold=True)
        C.put(ws, rr, 2, etf)
        C.put(ws, rr, 3, sn.get("cycle_fit_override") or fit, inp=True)
        C.put(ws, rr, 4, _yes_no(rv.get("direction", "")), inp=True)
        C.put(ws, rr, 5, "Yes" if bool(rs.get("rs_pillar")) else "No", inp=True)
        C.put(ws, rr, 6, _yes_no(va.get("supportive", "")), inp=True)
        C.put(ws, rr, 7, f"=COUNTIF(C{rr}:F{rr},\"Yes\")", "0")
        C.put(ws, rr, 8, f"=IF(G{rr}>=3,\"ENTER\",IF(G{rr}=2,\"WATCH\",\"AVOID\"))")
        C.put(ws, rr, 9, " | ".join(str(x) for x in (rv.get("value"), rv.get("source")) if x), wrap=True, ai=bool(rv))
        C.put(ws, rr, 10, " | ".join(str(x) for x in (va.get("value"), va.get("source")) if x), wrap=True, ai=bool(va))
        C.put(ws, rr, 11, sn.get("priced_in", ""), wrap=True, ai=bool(sn.get("priced_in")))
        C.put(ws, rr, 12, alloc.get(sector), "0", inp=True)
        C.put(ws, rr, 13, f"=IF(N(L{rr})=0,\"\",ROUND(L{rr}/100*{n_target},0))", "0")
        C.put(ws, rr, 14, ", ".join(subs.get(sector, [])), wrap=True, inp=True)
    last = first + len(SECTOR_ETFS) - 1
    C.put(ws, last + 1, 11, "Total", bold=True)
    C.put(ws, last + 1, 12, f"=SUM(L{first}:L{last})", "0", bold=True)
    C.put(ws, last + 1, 13, f"=SUM(M{first}:M{last})", "0", bold=True)
    C.put(ws, last + 2, 11, f"=IF(L{last + 1}=100,\"OK\",\"Weights should sum to 100\")")
    max_sec = ws_cfg.settings["portfolio"]["max_sector_pct"]
    C.put(ws, last + 3, 11, f"=IF(MAX(L{first}:L{last})>{max_sec},\"A sector exceeds the {max_sec}% cap\",\"Sector caps OK\")")
    C.status_colors(ws, f"C{first}:H{last}")
    for col, w in zip("ABCDEFGHIJKLMN", (24, 7, 10, 12, 12, 14, 10, 10, 40, 44, 34, 10, 12, 40)):
        ws.column_dimensions[col].width = w
    C.freeze(ws, f"C{first}")
    r = last + 5
    r = C.section(ws, r, "Entry and exit discipline for each sector call")
    for line in ["Split thesis from entry: pillars 1, 2 and 4 = thesis (fundamental); pillar 3 = market confirmation.",
                 "Scale in: 1/3 when the thesis is in place, 1/3 when relative strength confirms, 1/3 on a pullback.",
                 "Thesis stop (write it now): e.g. sector EPS revisions turn negative, or the yield curve / Fed path moves against the cycle call.",
                 "Price stop: about 2-3x ATR below entry, sized so a stop-out costs 1-2% of the portfolio.",
                 "Priced in? If the sector already trades at a historically high premium and everyone talks about the rotation, the remaining 6-month upside is smaller."]:
        C.put(ws, r, 1, line)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        r += 1

    # Correlation
    if corr is not None:
        ws = wb.create_sheet("Correlation")
        r = C.title(ws, "52-week correlation of weekly returns",
                    "Diversification check: two sectors correlated above ~0.85 behave like one bet.")
        cc = corr.drop(columns=["as_of"])
        C.table(ws, cc, r, 1, {k: "0.00" for k in cc.columns if k != "ticker"})
        n = len(cc)
        from openpyxl.utils import get_column_letter as L
        C.heat(ws, f"B{r + 1}:{L(n + 1)}{r + n}", reverse=True)

    # Events / Fed path / Sources from notes
    ws = wb.create_sheet("Events & Fed")
    r = C.title(ws, "Event calendar and policy path inside the horizon", "FOMC, CPI, jobs, earnings season. Sources required.")
    fp = notes.get("fed_path") or {}
    for lab, key in (("Market-implied path (CME FedWatch)", "market_implied"), ("Fed projections (dot plot / SEP)", "fed_dots"),
                     ("2Y minus fed funds (FRED proxy)", None)):
        C.put(ws, r, 1, lab, bold=True)
        val = fp.get(key) if key else f"{m['metrics'].get('two_year_minus_ff', float('nan')):+.2f} pp"
        C.put(ws, r, 2, val, wrap=True, ai=bool(key and fp.get(key)))
        r += 1
    r += 1
    ev = pd.DataFrame(notes.get("events") or [], columns=["date", "event", "why_it_matters", "source"])
    C.table(ws, ev, r, 1, wrap_cols=["why_it_matters"], widths={"date": 14, "event": 36, "why_it_matters": 70, "source": 50})
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 80

    ws = wb.create_sheet("Sources")
    r = C.title(ws, "Sources", "Primary sources first: Fed, BLS, BEA, FRED, FactSet Earnings Insight, CME FedWatch.")
    src = [{"title": f"FRED {k}", "url": v} for k, v in m["sources"].items()]
    src += [{"title": s.get("title"), "url": s.get("url"), "retrieved": s.get("retrieved")} for s in notes.get("sources") or []]
    C.table(ws, pd.DataFrame(src, columns=["title", "url", "retrieved"]), r, 1, widths={"title": 60, "url": 90, "retrieved": 14})

    C.readme(wb, "Macro, cycle and sector rotation", [
        ("Purpose", "Decide the cycle phase and which sectors deserve capital for the next 6 months."),
        ("Order", "Dashboard -> Cycle Rules (your call) -> Rotation -> Scorecard (3 of 4 pillars) -> weights and sub-industries."),
        ("Timing risk", "Markets price turns 6-9 months ahead. Requiring thesis + price confirmation trims false starts at the cost of "
                        "the first few percent of the move."),
        ("Next step", "Weights and sub-industries feed the sector-screening skill (stored in config/state.yaml)."),
    ])
    order = ["README", "Scorecard", "Cycle Rules", "Dashboard", "Rotation", "Correlation", "Events & Fed", "Sources"]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames]
    out = ws_cfg.output / "macro" / f"Macro_{store.today()}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
